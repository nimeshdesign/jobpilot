"""Render a tailored resume + cover letter to DOCX and PDF.

Layout is deliberately plain: single column, standard fonts, real headings, no
tables, no text boxes, no graphics. That is what ATS parsers can actually read —
the pretty two-column template is the reason many resumes come out as garbage on
the employer's side.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from ..config import GENERATED_DIR


# ReportLab's built-in fonts render U+2022 ("•") as a glyph that text
# extraction cannot map back to a character -- it comes out as "(cid:127)".
# An ATS parses the PDF exactly the way that extractor does, so a pretty bullet
# it cannot read is worse than a plain one it can. U+00B7 survives extraction.
PDF_BULLET = "·"
SKILL_SEPARATOR = f" {PDF_BULLET} "


def _slug(text: str, limit: int = 40) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").strip()).strip("-").lower()
    return (slug or "untitled")[:limit]


def output_basename(job, application_id: int) -> str:
    return f"{application_id:05d}-{_slug(job.company)}-{_slug(job.title, 30)}"


def _contact_line(contact: dict, profile) -> str:
    parts = [
        contact.get("email") or profile.email,
        contact.get("phone") or profile.phone,
        contact.get("location") or ", ".join(filter(None, [profile.city, profile.country])),
        contact.get("linkedin") or profile.linkedin,
        contact.get("github") or profile.github,
        contact.get("portfolio") or profile.portfolio,
    ]
    return "  |  ".join(p for p in parts if p)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def write_docx(tailored, resume, profile, path: Path) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    parsed = resume.parsed or {}
    contact = parsed.get("contact", {})

    document = docx.Document()

    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(40)
        section.left_margin = section.right_margin = Pt(50)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def heading(text: str) -> None:
        para = document.add_paragraph()
        para.paragraph_format.space_before = Pt(11)
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # --- header ---
    name_para = document.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_para.paragraph_format.space_after = Pt(2)
    name_run = name_para.add_run(contact.get("name") or profile.full_name or "Your Name")
    name_run.bold = True
    name_run.font.size = Pt(19)

    if tailored.headline:
        head_para = document.add_paragraph()
        head_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        head_para.paragraph_format.space_after = Pt(2)
        head_run = head_para.add_run(tailored.headline)
        head_run.font.size = Pt(11)
        head_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    contact_para = document.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.paragraph_format.space_after = Pt(6)
    contact_run = contact_para.add_run(_contact_line(contact, profile))
    contact_run.font.size = Pt(9)

    # --- summary ---
    if tailored.summary:
        heading("Professional Summary")
        document.add_paragraph(tailored.summary)

    # --- skills ---
    if tailored.top_skills:
        heading("Core Skills")
        document.add_paragraph(" • ".join(tailored.top_skills))

    # --- experience ---
    if tailored.experience:
        heading("Professional Experience")
        for role in tailored.experience:
            role_para = document.add_paragraph()
            role_para.paragraph_format.space_before = Pt(6)
            role_para.paragraph_format.space_after = Pt(1)
            title_run = role_para.add_run(role.title or "")
            title_run.bold = True
            if role.company:
                role_para.add_run(f" — {role.company}")
            if role.dates:
                date_run = role_para.add_run(f"    ({role.dates})")
                date_run.font.size = Pt(9.5)
                date_run.italic = True
            for bullet in role.bullets:
                bullet_para = document.add_paragraph(bullet, style="List Bullet")
                bullet_para.paragraph_format.space_after = Pt(1)

    # --- education / certs / projects come straight from the original resume ---
    education = parsed.get("education") or []
    if education:
        heading("Education")
        for entry in education[:4]:
            document.add_paragraph(entry.get("text", ""))

    certifications = parsed.get("certifications") or []
    if certifications:
        heading("Certifications")
        for cert in certifications[:6]:
            document.add_paragraph(cert, style="List Bullet")

    projects = parsed.get("projects") or []
    if projects:
        heading("Projects")
        for project in projects[:5]:
            document.add_paragraph(project, style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def write_pdf(tailored, resume, profile, path: Path) -> Path:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    parsed = resume.parsed or {}
    contact = parsed.get("contact", {})
    base = getSampleStyleSheet()

    style_name = ParagraphStyle("Name", parent=base["Title"], fontSize=19, leading=23,
                                spaceAfter=2, alignment=TA_CENTER)
    style_headline = ParagraphStyle("Headline", parent=base["Normal"], fontSize=10.5,
                                    alignment=TA_CENTER, textColor="#444444", spaceAfter=2)
    style_contact = ParagraphStyle("Contact", parent=base["Normal"], fontSize=8.5,
                                   alignment=TA_CENTER, spaceAfter=10)
    style_section = ParagraphStyle("Section", parent=base["Heading2"], fontSize=11, leading=13,
                                   spaceBefore=11, spaceAfter=3, textColor="#1A365D")
    style_body = ParagraphStyle("Body", parent=base["Normal"], fontSize=10, leading=13.5,
                                spaceAfter=3)
    style_role = ParagraphStyle("Role", parent=base["Normal"], fontSize=10.5, leading=13,
                                spaceBefore=6, spaceAfter=1)
    style_bullet = ParagraphStyle("Bullet", parent=base["Normal"], fontSize=9.8, leading=12.5)

    def esc(text: str) -> str:
        return (
            (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )

    story = [
        Paragraph(esc(contact.get("name") or profile.full_name or "Your Name"), style_name)
    ]
    if tailored.headline:
        story.append(Paragraph(esc(tailored.headline), style_headline))
    story.append(Paragraph(esc(_contact_line(contact, profile)), style_contact))

    def section(title: str) -> None:
        story.append(Paragraph(esc(title.upper()), style_section))

    if tailored.summary:
        section("Professional Summary")
        story.append(Paragraph(esc(tailored.summary), style_body))

    if tailored.top_skills:
        section("Core Skills")
        story.append(Paragraph(
            SKILL_SEPARATOR.join(esc(skill) for skill in tailored.top_skills), style_body
        ))

    if tailored.experience:
        section("Professional Experience")
        for role in tailored.experience:
            header = f"<b>{esc(role.title)}</b>"
            if role.company:
                header += f" — {esc(role.company)}"
            if role.dates:
                header += f" <font size=8.5><i>({esc(role.dates)})</i></font>"
            story.append(Paragraph(header, style_role))
            if role.bullets:
                story.append(ListFlowable(
                    [ListItem(Paragraph(esc(b), style_bullet), leftIndent=12)
                     for b in role.bullets],
                    bulletType="bullet", start=PDF_BULLET, leftIndent=14,
                ))

    education = parsed.get("education") or []
    if education:
        section("Education")
        for entry in education[:4]:
            story.append(Paragraph(esc(entry.get("text", "")), style_body))

    certifications = parsed.get("certifications") or []
    if certifications:
        section("Certifications")
        story.append(ListFlowable(
            [ListItem(Paragraph(esc(c), style_bullet)) for c in certifications[:6]],
            bulletType="bullet", start=PDF_BULLET, leftIndent=14,
        ))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path), pagesize=LETTER,
        leftMargin=0.7 * inch, rightMargin=0.7 * inch,
        topMargin=0.55 * inch, bottomMargin=0.55 * inch,
        title=f"{contact.get('name') or profile.full_name} — Resume",
        author=contact.get("name") or profile.full_name,
    )
    doc.build(story)
    return path


def write_cover_letter_pdf(text: str, tailored, resume, profile, job, path: Path) -> Path:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    parsed = resume.parsed or {}
    contact = parsed.get("contact", {})
    base = getSampleStyleSheet()
    style_body = ParagraphStyle("Body", parent=base["Normal"], fontSize=10.5, leading=15,
                                spaceAfter=9)
    style_meta = ParagraphStyle("Meta", parent=base["Normal"], fontSize=9, leading=12,
                                spaceAfter=2)

    def esc(value: str) -> str:
        return (value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [
        Paragraph(f"<b>{esc(contact.get('name') or profile.full_name)}</b>", style_meta),
        Paragraph(esc(_contact_line(contact, profile)), style_meta),
        Spacer(1, 16),
        Paragraph(esc(date.today().strftime("%d %B %Y")), style_meta),
        Spacer(1, 10),
        Paragraph(f"<b>Re: {esc(job.title)}{' at ' + esc(job.company) if job.company else ''}</b>",
                  style_body),
        Spacer(1, 4),
    ]
    for para in (text or "").split("\n\n"):
        cleaned = para.strip()
        if cleaned:
            story.append(Paragraph(esc(cleaned).replace("\n", "<br/>"), style_body))

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path), pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.8 * inch, bottomMargin=0.8 * inch,
        title=f"Cover letter — {job.title}",
    ).build(story)
    return path


def purge_previous(application_id: int, keep_prefix: str) -> int:
    """Delete documents left over from an earlier run of this application.

    Filenames embed the company and job title, so re-preparing an application
    against a different job writes a new file and orphans the old one. Both then
    sit in the folder under the same "00007-" prefix and it is pure luck which
    one you open -- which is how a cover letter for the wrong company gets sent.
    """
    removed = 0
    for path in GENERATED_DIR.glob(f"{application_id:05d}-*"):
        if path.name.startswith(keep_prefix):
            continue
        try:
            path.unlink()
            removed += 1
        except OSError:
            pass
    return removed


def generate_all(tailored, resume, profile, job, application_id: int) -> dict[str, str]:
    """Write resume DOCX + PDF and (optionally) a cover-letter PDF."""
    base = output_basename(job, application_id)
    purge_previous(application_id, base)
    out: dict[str, str] = {}

    docx_path = GENERATED_DIR / f"{base}-resume.docx"
    pdf_path = GENERATED_DIR / f"{base}-resume.pdf"

    write_docx(tailored, resume, profile, docx_path)
    out["docx"] = str(docx_path)

    try:
        write_pdf(tailored, resume, profile, pdf_path)
        out["pdf"] = str(pdf_path)
    except Exception as exc:  # noqa: BLE001 - a missing PDF must not block applying
        out["pdf_error"] = f"{type(exc).__name__}: {exc}"

    if tailored.cover_letter:
        letter_path = GENERATED_DIR / f"{base}-cover-letter.pdf"
        try:
            write_cover_letter_pdf(
                tailored.cover_letter, tailored, resume, profile, job, letter_path
            )
            out["cover_letter_pdf"] = str(letter_path)
        except Exception as exc:  # noqa: BLE001
            out["cover_letter_error"] = f"{type(exc).__name__}: {exc}"

    return out
